import io
import logging

logger = logging.getLogger("crypto_watchman.document_parser")

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}
MAX_DOCUMENT_BYTES = 10 * 1024 * 1024  # 10 MB


def extract_document_extension(filename: str) -> str:
    return f".{filename.rsplit('.', 1)[-1].lower()}" if "." in filename else ""


def extract_text_from_document(file_bytes: bytes, filename: str) -> str:
    """Extract readable text from a PDF, DOCX or plain-text strategy document."""
    ext = extract_document_extension(filename)

    if ext == ".pdf":
        return _extract_pdf(file_bytes)
    if ext == ".docx":
        return _extract_docx(file_bytes)
    if ext in (".txt", ".md"):
        return _extract_plain_text(file_bytes)
    raise ValueError(
        f"Unsupported file type '{ext or 'unknown'}'. Please upload a PDF, DOCX, TXT or MD document."
    )


def _extract_pdf(file_bytes: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError:
        raise ValueError("PDF parsing is unavailable (pypdf not installed).")

    reader = PdfReader(io.BytesIO(file_bytes))
    pages = []
    for page in reader.pages[:40]:
        text = page.extract_text() or ""
        pages.append(text)
    return "\n".join(pages).strip()


def _extract_docx(file_bytes: bytes) -> str:
    try:
        from docx import Document
    except ImportError:
        raise ValueError("DOCX parsing is unavailable (python-docx not installed).")

    doc = Document(io.BytesIO(file_bytes))
    parts = []
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text.strip())
    for table in doc.tables[:20]:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts).strip()


def _extract_plain_text(file_bytes: bytes) -> str:
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            return file_bytes.decode(encoding).strip()
        except (UnicodeDecodeError, LookupError):
            continue
    return file_bytes.decode("utf-8", errors="replace").strip()